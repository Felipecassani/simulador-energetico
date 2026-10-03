"""Consumos de 15 em 15 minutos do Balcão Digital da E-REDES (Excel ou CSV).

Formato confirmado com um ficheiro real (01/10/2026), folha "Dados de Energia":
  linhas de cabeçalho (CPE, "Mês/Ano", "Intervalo: 15 min"), depois
  Data | Hora | Consumo registado (kW) | Estado     ex.: 2026/09/01 | 00:15 | 0,464 | Real
  A hora marca o FIM do quarto de hora; o último intervalo de cada dia vem como 00:00 do dia
  seguinte. O consumo é a potência média (kW) do quarto de hora → kWh = kW × 0,25.
A leitura continua tolerante a variações:
  - procura a linha de cabeçalho com "Data" e "Hora" (ou uma coluna de data e hora juntas);
  - a coluna de consumo é a que tem "consumo", "kw" ou "energia" no nome;
  - em kW (potência média do quarto de hora) multiplica por 0,25 h; em kWh usa direto;
  - se o primeiro registo é 00:15 (ou aparece 24:00), a hora marca o fim do intervalo.
"""
import csv
import io
import re
import zipfile
from itertools import islice
from collections import defaultdict
from datetime import date, datetime, time, timedelta

from nucleo import periodos

QUARTO = timedelta(minutes=15)
# Proteção do servidor: um .xlsx é um ZIP e pode expandir-se muito ("zip bomb"). Um ano da E-REDES
# tem ~35 000 linhas e poucos MB descomprimido.
MAX_DESCOMPRIMIDO = 60 * 1024 * 1024        # bytes, somando todos os ficheiros dentro do .xlsx
MAX_LINHAS = 60_000                         # mais de um ano e meio de quartos de hora


def _celulas_xlsx(conteudo):
    import warnings

    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_DESCOMPRIMIDO:
            raise ValueError("O ficheiro é demasiado grande depois de descomprimido.")
    with warnings.catch_warnings():             # o ficheiro da E-REDES não tem estilo por omissão
        warnings.simplefilter("ignore")
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    for folha in livro.worksheets:
        yield from folha.iter_rows(values_only=True)


def _celulas_csv(conteudo):
    texto = conteudo.decode("utf-8-sig", errors="replace")
    separador = ";" if texto.count(";") >= texto.count(",") else ","
    yield from csv.reader(io.StringIO(texto), delimiter=separador)


def _numero(v):
    if isinstance(v, (int, float)):
        return float(v)
    texto = str(v or "").strip().replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _data(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    texto = str(v or "").strip()
    for formato in ("%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M",
                    "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def _hora(v):
    """(horas, minutos) — aceita time, datetime, '00:15', '24:00'."""
    if isinstance(v, (time, datetime)):
        return v.hour, v.minute
    m = re.match(r"^\s*(\d{1,2})[:h](\d{2})", str(v or ""))
    return (int(m.group(1)), int(m.group(2))) if m else None


def ler(nome, conteudo):
    """→ [(início do quarto de hora em Lisboa, kWh)] por ordem."""
    return ler_com_estimados(nome, conteudo)[0]


def ler_com_estimados(nome, conteudo):
    """→ (registos, quantos valores eram "Estimado") — sem estado partilhado entre sessões."""
    celulas = _celulas_xlsx(conteudo) if nome.lower().endswith((".xlsx", ".xlsm")) else _celulas_csv(conteudo)
    linhas = list(islice(celulas, MAX_LINHAS + 1))
    if len(linhas) > MAX_LINHAS:
        raise ValueError("O ficheiro tem mais linhas do que um ano e meio de consumos.")
    cabecalho, i_cab = None, None
    for i, linha in enumerate(linhas):
        nomes = [str(c or "").strip().lower() for c in linha]
        if any("data" in n for n in nomes) and any(("consumo" in n or "kw" in n or "energia" in n)
                                                     for n in nomes):
            cabecalho, i_cab = nomes, i
            break
    if cabecalho is None:
        raise ValueError("Não encontrei as colunas de data e consumo neste ficheiro.")
    c_data = next(i for i, n in enumerate(cabecalho) if "data" in n)
    c_hora = next((i for i, n in enumerate(cabecalho) if n.startswith("hora")), None)
    c_valor = next(i for i, n in enumerate(cabecalho)
                   if ("consumo" in n or "kw" in n or "energia" in n) and i not in (c_data, c_hora))
    em_kw = "kw" in cabecalho[c_valor] and "kwh" not in cabecalho[c_valor]
    c_estado = next((i for i, n in enumerate(cabecalho) if n.startswith("estado")), None)

    brutos = []
    for linha in linhas[i_cab + 1:]:
        if len(linha) <= max(c_data, c_valor, c_hora or 0):
            continue
        dia, valor = _data(linha[c_data]), _numero(linha[c_valor])
        if dia is None or valor is None:
            continue
        hm = _hora(linha[c_hora]) if c_hora is not None else (dia.hour, dia.minute)
        if hm is None:
            continue
        estimado = c_estado is not None and len(linha) > c_estado and "estim" in str(linha[c_estado] or "").lower()
        brutos.append((dia.replace(hour=0, minute=0, second=0), hm, valor * 0.25 if em_kw else valor, estimado))
    if not brutos:
        raise ValueError("O ficheiro não tem consumos de 15 em 15 minutos.")
    # a E-REDES marca o fim do intervalo: o 1.º registo é 00:15 e o último de cada dia é 00:00 do seguinte.
    # A coluna "Consumo registado" é da E-REDES: fim do intervalo mesmo que o ficheiro comece a meio do dia.
    marca_fim = (brutos[0][1] == (0, 15) or any(hm == (24, 0) for _, hm, _, _ in brutos)
                 or "consumo registado" in cabecalho[c_valor])
    registos = []
    for dia, (h, m), kwh, _ in brutos:
        inicio = dia + timedelta(hours=h, minutes=m) - (QUARTO if marca_fim else timedelta(0))
        registos.append((inicio.replace(tzinfo=periodos.LISBOA), kwh))
    return sorted(registos), sum(1 for *_, e in brutos if e)


def analisar(registos):
    """Totais, repartição pelos períodos, pico e consumo médio por hora do dia."""
    total = sum(k for _, k in registos)
    dias = max(1, round(len(registos) / 96))            # dias completos de quartos de hora
    por_periodo = {o: defaultdict(float) for o in ("bi", "tri")}
    por_hora = defaultdict(float)
    for t, kwh in registos:
        for o in por_periodo:
            por_periodo[o][periodos.periodo(t, o)] += kwh
        por_hora[t.hour] += kwh
    pct = lambda o, p: 100 * por_periodo[o][p] / total if total else 0.0
    return {
        "total_kwh": total, "dias": dias, "inicio": registos[0][0].date(), "fim": registos[-1][0].date(),
        "pct_vazio": pct("tri", "vazio"), "pct_ponta": pct("tri", "ponta"),
        "pico_kw": max(k for _, k in registos) * 4,
        "media_por_hora": [por_hora[h] / dias if dias else 0.0 for h in range(24)],
    }


def preco_ponderado(registos, precos_omie):
    """€/MWh médio que o teu perfil paga no mercado (ponderado pelo consumo) e a média simples.

    precos_omie: [(instante, €/MWh)]. Quartos de hora sem preço ficam de fora.
    """
    mapa = {t.astimezone(periodos.LISBOA).replace(second=0, microsecond=0): v for t, v in precos_omie}
    horario = len(precos_omie) > 1 and (precos_omie[1][0] - precos_omie[0][0]) >= timedelta(hours=1)
    soma, kwh_total, valores = 0.0, 0.0, []
    for t, kwh in registos:
        chave = t.replace(minute=0) if horario else t
        if chave in mapa:
            soma += kwh * mapa[chave]
            kwh_total += kwh
            valores.append(mapa[chave])
    if not kwh_total:
        return None
    return {"ponderado": soma / kwh_total, "simples": sum(valores) / len(valores)}


DIAS_SEMANA = ("Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom")


def padroes(registos):
    """Padrões de um período longo (até um ano): por mês, por hora × mês, dias da semana e consumo de base.

    - meses: kWh, kWh/dia, % vazio e ponta (tri-horário) e pico (kW) de cada mês;
    - mapa: consumo médio por dia em cada hora, para cada mês (o "mapa de calor");
    - dias úteis contra fim de semana (só dias completos, com os 96 quartos de hora);
    - consumo de base: o que fica sempre ligado (frigorífico, standby, routers), medido como o
      percentil 10 da potência entre as 2h e as 5h, quando quase tudo está desligado.
    """
    por_mes = defaultdict(lambda: {"kwh": 0.0, "vazio": 0.0, "ponta": 0.0, "pico_kw": 0.0, "quartos": 0})
    hora_mes, dias_mes = defaultdict(float), defaultdict(set)
    por_dia, quartos_dia, noite = defaultdict(float), defaultdict(int), []
    for t, kwh in registos:
        m = (t.year, t.month)
        d = por_mes[m]
        d["kwh"] += kwh
        d["quartos"] += 1
        periodo = periodos.periodo(t, "tri")
        if periodo in ("vazio", "ponta"):
            d[periodo] += kwh
        d["pico_kw"] = max(d["pico_kw"], kwh * 4)
        hora_mes[(m, t.hour)] += kwh
        dias_mes[m].add(t.date())
        por_dia[t.date()] += kwh
        quartos_dia[t.date()] += 1
        if 2 <= t.hour < 5:
            noite.append(kwh * 4)

    meses = []
    for m in sorted(por_mes):
        d = por_mes[m]
        n_dias = d["quartos"] / 96
        meses.append({"mes": date(m[0], m[1], 1), "kwh": d["kwh"], "dias": n_dias,
                      "kwh_dia": d["kwh"] / n_dias if n_dias else 0.0,
                      "pct_vazio": 100 * d["vazio"] / d["kwh"] if d["kwh"] else 0.0,
                      "pct_ponta": 100 * d["ponta"] / d["kwh"] if d["kwh"] else 0.0,
                      "pico_kw": d["pico_kw"]})
    # dividido pelos mesmos dias do kWh/dia do mês (quartos/96): a soma das 24 horas bate com ele
    mapa = {date(m[0], m[1], 1): [hora_mes[(m, h)] / (por_mes[m]["quartos"] / 96) for h in range(24)]
            for m in sorted(por_mes)}

    completos = {dia: v for dia, v in por_dia.items() if quartos_dia[dia] >= 92}   # 23 h no dia da mudança de hora
    semana = defaultdict(list)
    for dia, v in completos.items():
        semana[dia.weekday()].append(v)
    media = lambda vs: sum(vs) / len(vs) if vs else None
    uteis = [v for dia, v in completos.items() if dia.weekday() < 5]
    fds = [v for dia, v in completos.items() if dia.weekday() >= 5]
    noite.sort()
    base_kw = noite[len(noite) // 10] if noite else None
    maior = max(completos.items(), key=lambda x: x[1]) if completos else None
    return {
        "meses": meses, "mapa": mapa,
        "por_dia_semana": [media(semana[i]) for i in range(7)],
        "kwh_dia_util": media(uteis), "kwh_dia_fds": media(fds),
        "base_kw": base_kw, "base_kwh_ano": base_kw * 24 * 365 if base_kw is not None else None,
        "dia_maior": maior, "dias_completos": len(completos),
    }
